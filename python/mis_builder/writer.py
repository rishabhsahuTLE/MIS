"""Write/replace the 'Working' and 'Sheet1' sheets inside an openpyxl
Workbook, given the rows built by working.py and sheet1.py.
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties

from .config import SHEET1_HEADERS, SHEET1_NAME, WORKING_HEADERS, WORKING_SHEET_NAME
from .sheet1 import Sheet1Layout
from .working import WorkingRow

HEADER_FONT = Font(bold=True)
DATE_FORMAT = "dd-mm-yyyy"


def _replace_sheet(wb: Workbook, name: str):
    if name in wb.sheetnames:
        del wb[name]
    return wb.create_sheet(title=name)


def write_working_sheet(wb: Workbook, rows: list[WorkingRow]):
    ws = _replace_sheet(wb, WORKING_SHEET_NAME)
    ws.append(WORKING_HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for w in rows:
        ws.append([
            w.invoice_or_credit_note,   # A
            w.travel_month,             # B
            w.name_of_passenger,        # C
            w.location,                 # D
            w.cost_code,                # E
            w.emp_id,                   # F
            w.nature_of_transaction,    # G
            w.nature_of_ticket,         # H
            w.amounts,                  # I
            w.invoice_number,           # J
            w.invoice_date,             # K
            w.gl,                       # L
            w.gl_code_formula,          # M
            w.narration_formula,        # N
            w.vendor_id,                # O
            w.submission_date,          # P
            w.gl_code_dup_formula,      # Q
            w.bu_formula,               # R
            w.service_month,            # S
        ])
        if w.invoice_date is not None:
            ws.cell(row=ws.max_row, column=11).number_format = DATE_FORMAT
    return ws


def _write_sheet1_row(ws, row_num: int, row):
    ws.cell(row=row_num, column=1, value=row.name_of_passenger)
    ws.cell(row=row_num, column=2, value=row.emp_id)
    ws.cell(row=row_num, column=3, value=row.invoice_number)
    date_cell = ws.cell(row=row_num, column=4, value=row.invoice_date)
    if row.invoice_date is not None:
        date_cell.number_format = DATE_FORMAT
    ws.cell(row=row_num, column=5, value=row.gl_code)
    ws.cell(row=row_num, column=6, value=row.narration)
    ws.cell(row=row_num, column=7, value=row.vendor_id)
    ws.cell(row=row_num, column=8, value=row.invoice_or_credit_note)
    ws.cell(row=row_num, column=9, value=row.amounts)


def write_sheet1(wb: Workbook, layout: Sheet1Layout):
    ws = _replace_sheet(wb, SHEET1_NAME)
    for col, header in enumerate(SHEET1_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT

    # Table 1 ('Invoice'): helper columns K (GL code, static) and L (GL
    # suffix, static) after a blank spacer column J.
    row_num = layout.table1_start_row
    for row in layout.table1:
        _write_sheet1_row(ws, row_num, row)
        ws.cell(row=row_num, column=11, value=row.gl_code)
        ws.cell(row=row_num, column=12, value=row.cost_code)
        row_num += 1
    if layout.table1:
        ws.cell(
            row=layout.table1_total_row, column=9,
            value=f"=SUM(I{layout.table1_start_row}:I{layout.table1_total_row - 1})",
        )

    # Table 2 ('Credit Note'): no repeated header. Helper columns J (literal
    # '-'), K (=CONCATENATE(J,I) -> negative amount), L (GL code, static),
    # M (GL suffix, static).
    row_num = layout.table2_start_row
    for row in layout.table2:
        _write_sheet1_row(ws, row_num, row)
        ws.cell(row=row_num, column=10, value="-")
        ws.cell(row=row_num, column=11, value=f"=CONCATENATE(J{row_num},I{row_num})")
        ws.cell(row=row_num, column=12, value=row.gl_code)
        ws.cell(row=row_num, column=13, value=row.cost_code)
        row_num += 1
    if layout.table2:
        ws.cell(
            row=layout.table2_total_row, column=9,
            value=f"=SUM(I{layout.table2_start_row}:I{layout.table2_total_row - 1})",
        )
    return ws


def build_output_workbook(wb: Workbook, working_rows: list[WorkingRow], sheet1_layout: Sheet1Layout):
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    write_working_sheet(wb, working_rows)
    write_sheet1(wb, sheet1_layout)
    return wb
