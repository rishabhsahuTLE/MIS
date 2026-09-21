"""Command-line entry point.

Usage:
    python -m mis_builder.cli --input "raw MIS.xlsx" [--output built.xlsx]
        [--travel-month AUG-26] [--service-month AUG-26]
        [--submission-date 06-05-2026] [--gl-prefix 80003]
        [--vendor-id V03136]
"""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

from .config import PeriodConfig
from .pipeline import SheetLayoutError, process_workbook


def _default_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_processed{input_path.suffix}")


def build_arg_parser() -> argparse.ArgumentParser:
    defaults = PeriodConfig()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Path to the raw MIS workbook.")
    parser.add_argument("--output", help="Path to write the result to (default: <input>_processed.xlsx).")
    parser.add_argument("--travel-month", default=defaults.travel_month)
    parser.add_argument("--service-month", default=defaults.service_month)
    parser.add_argument("--submission-date", default=defaults.submission_date)
    parser.add_argument("--gl-prefix", default=defaults.gl_prefix)
    parser.add_argument("--vendor-id", default=defaults.vendor_id)
    return parser


def run(args=None) -> int:
    parser = build_arg_parser()
    opts = parser.parse_args(args)

    input_path = Path(opts.input)
    output_path = Path(opts.output) if opts.output else _default_output_path(input_path)

    period = PeriodConfig(
        travel_month=opts.travel_month,
        service_month=opts.service_month,
        submission_date=opts.submission_date,
        gl_prefix=opts.gl_prefix,
        vendor_id=opts.vendor_id,
    )

    wb = load_workbook(input_path, data_only=False)
    wb_values = load_workbook(input_path, data_only=True)

    print(f"Reading source sheets from {input_path.name}:")
    try:
        result = process_workbook(wb, wb_values, period)
    except SheetLayoutError as exc:
        print(f"  ! {exc}", file=sys.stderr)
        raise

    for name, count in result.rows_per_sheet.items():
        print(f"  - {name}: {count} rows")
    for name in result.sheets_not_found:
        print(f"  - {name}: NOT FOUND, skipping")
    for warning in result.unparsed_invoice_dates:
        print(f"  ! {warning}")

    wb.save(output_path)

    print()
    print(f"Working: {result.working_row_count} rows written.")
    print(
        f"Sheet1: {result.sheet1_invoice_count} Invoice rows, "
        f"{result.sheet1_credit_note_count} Credit Note rows."
    )
    print(f"Saved to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(run())
